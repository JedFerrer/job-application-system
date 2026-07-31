from __future__ import annotations

from collections.abc import Callable
from io import BytesIO

import pytest
from docx import Document
from docx.document import Document as DocxDocument

from app.infrastructure.resume_parser import extract_text_from_bytes

pytestmark = pytest.mark.unit


def _docx_bytes(build: Callable[[DocxDocument], object]) -> bytes:
    document = Document()
    build(document)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_paragraph_only_docx_extraction_is_unchanged() -> None:
    content = _docx_bytes(
        lambda document: (
            document.add_paragraph("First paragraph"),
            document.add_paragraph("Second paragraph"),
        )
    )

    assert extract_text_from_bytes(content, "resume.docx") == (
        "First paragraph\n\nSecond paragraph"
    )


def test_table_only_docx_extraction_succeeds() -> None:
    def build(document: DocxDocument) -> None:
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Skill"
        table.cell(0, 1).text = "Python"
        table.cell(1, 0).text = "Project"
        table.cell(1, 1).text = "Resume matcher"

    assert extract_text_from_bytes(_docx_bytes(build), "resume.docx") == (
        "Skill | Python\n\nProject | Resume matcher"
    )


def test_paragraphs_and_tables_keep_source_order() -> None:
    def build(document: DocxDocument) -> None:
        document.add_paragraph("Before table")
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Inside table"
        document.add_paragraph("After table")

    assert extract_text_from_bytes(_docx_bytes(build), "resume.docx") == (
        "Before table\n\nInside table\n\nAfter table"
    )


def test_empty_table_cells_are_skipped_and_rows_stay_ordered() -> None:
    def build(document: DocxDocument) -> None:
        table = document.add_table(rows=2, cols=3)
        table.cell(0, 0).text = ""
        table.cell(0, 1).text = "Python"
        table.cell(0, 2).text = "  "
        table.cell(1, 0).text = "Project"
        table.cell(1, 1).text = "Matcher"
        table.cell(1, 2).text = ""

    assert extract_text_from_bytes(_docx_bytes(build), "resume.docx") == (
        "Python\n\nProject | Matcher"
    )


def test_table_cells_with_multiple_paragraphs_are_flattened_safely() -> None:
    def build(document: DocxDocument) -> None:
        table = document.add_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        cell.text = "First line"
        cell.add_paragraph("Second line")

    assert extract_text_from_bytes(_docx_bytes(build), "resume.docx") == ("First line Second line")


def test_empty_docx_preserves_actionable_error() -> None:
    with pytest.raises(
        ValueError,
        match=r"No text could be extracted from DOCX. Is the document empty\?",
    ):
        extract_text_from_bytes(_docx_bytes(lambda _document: None), "resume.docx")
