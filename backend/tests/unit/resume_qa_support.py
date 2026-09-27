"""Test-only helpers for tailored-resume Phase 4 QA.

These utilities do not change production behavior. They materialize the
sanitized wrap fixture, build source-backed snapshots, and normalize
preview/copy/PDF/DOCX text for semantic comparison.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from collections import Counter
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from typing import Any, cast

from docx import Document

from app.domain.resume.schemas import (
    ParsedProfile,
    SourceResumeContentDocumentV2,
    SourceResumeContentDocumentV4,
    SourceResumeDocument,
    SourceResumeLayoutReferencesV4,
    TailoredResume,
    TailoredWorkItem,
)
from app.domain.resume.selection import select_source_bullets
from app.domain.resume.service import (
    _canonical_sections,
    _required_included_indices,
)
from app.domain.resume.source_content import (
    build_source_candidate_ledger,
    canonicalize_source_layout_v5,
    classify_source_artifacts,
    materialize_source_document_v4,
    resolve_source_layout_v4,
    source_revision_sha256,
    verify_source_candidate_ledger,
)

FIXTURE_DIR = Path(__file__).resolve().parents[1] / "fixtures" / "tailored_resume_qa"
WRAP_RESUME_PATH = FIXTURE_DIR / "malformed_wrap_two_role.txt"
EXPECTED_SEMANTIC_PATH = FIXTURE_DIR / "expected_semantic.json"

_WHITESPACE_RE = re.compile(r"\s+")


def load_wrap_resume_text() -> str:
    return WRAP_RESUME_PATH.read_text(encoding="utf-8")


def load_expected_semantic() -> dict[str, Any]:
    loaded = json.loads(EXPECTED_SEMANTIC_PATH.read_text(encoding="utf-8"))
    if not isinstance(loaded, dict):
        raise AssertionError("expected_semantic.json must be an object")
    return cast(dict[str, Any], loaded)


def normalize_resume_export_text(value: str) -> str:
    """Collapse formatting so preview/copy/PDF/DOCX can be compared semantically."""
    normalized = unicodedata.normalize("NFKC", value)
    normalized = normalized.replace("\u2022", " ")
    normalized = _WHITESPACE_RE.sub(" ", normalized)
    return normalized.strip()


def count_normalized(haystack: str, needle: str) -> int:
    haystack_n = normalize_resume_export_text(haystack)
    needle_n = normalize_resume_export_text(needle)
    if not needle_n:
        return 0
    count = 0
    start = 0
    while True:
        found = haystack_n.find(needle_n, start)
        if found < 0:
            return count
        count += 1
        start = found + len(needle_n)


def semantic_tokens(expected: dict[str, Any] | None = None) -> list[str]:
    payload = expected if expected is not None else load_expected_semantic()
    tokens: list[str] = []
    for role in payload["roles"]:
        tokens.append(str(role["header"]))
        narrative = str(role.get("narrative") or "").strip()
        if narrative:
            tokens.append(narrative)
        heading = str(role.get("responsibilities_heading") or "").strip()
        if heading:
            tokens.append(heading)
        tokens.extend(str(item) for item in role["bullets"])
    return tokens


def assert_surface_matches_expected(surface: str, expected: dict[str, Any] | None = None) -> None:
    """Require complete semantic tokens in order, with no duplicated fragments."""
    payload = expected if expected is not None else load_expected_semantic()
    tokens = semantic_tokens(payload)
    normalized = normalize_resume_export_text(surface)
    cursor = 0
    for token in tokens:
        needle = normalize_resume_export_text(token)
        found = normalized.find(needle, cursor)
        if found < 0:
            raise AssertionError(f"semantic token missing from surface: {token!r}")
        cursor = found + len(needle)
    counts = Counter(normalize_resume_export_text(item) for item in tokens)
    for token, expected_count in counts.items():
        actual = count_normalized(surface, token)
        if actual != expected_count:
            raise AssertionError(f"semantic token count {actual} != {expected_count}: {token!r}")
    for fragment in payload["orphan_fragments"]:
        if count_normalized(surface, str(fragment)) != 1:
            raise AssertionError(
                f"orphan continuation fragment is not confined to its parent token: {fragment!r}"
            )


def extract_docx_text(content: bytes) -> str:
    document = Document(BytesIO(content))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def extract_pdf_text(content: bytes) -> str:
    import pdfplumber

    with pdfplumber.open(BytesIO(content)) as pdf:
        pages = [page.extract_text() or "" for page in pdf.pages]
    return "\n".join(pages)


def write_semantic_docx(expected: dict[str, Any] | None = None) -> bytes:
    """Generate a comparison DOCX from expected semantic tokens (test-only)."""
    document = Document()
    for token in semantic_tokens(expected):
        document.add_paragraph(token)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _alias_for(ledger: SourceResumeContentDocumentV2, text: str) -> str:
    for candidate in ledger.candidates:
        if candidate.text == text:
            return f"c{candidate.order}"
    raise AssertionError(f"missing source candidate: {text}")


def _empty_v5_layout(*headers: str) -> dict[str, object]:
    return {
        "contact": {},
        "links": {},
        "summary": None,
        "summary_lines": [],
        "section_headings": {},
        "skills": [],
        "sections": [],
        "work_history": [
            {
                "order": index,
                "header": header,
                "context": [],
                "responsibilities_heading": None,
                "highlights": [],
            }
            for index, header in enumerate(headers)
        ],
        "projects": [],
        "education": [],
        "certifications": [],
        "licenses": [],
        "languages": [],
        "named_sections": [],
    }


def materialize_wrap_source_document(
    raw: str | None = None,
) -> tuple[SourceResumeContentDocumentV2, SourceResumeDocument]:
    text = raw if raw is not None else load_wrap_resume_text()
    expected = load_expected_semantic()
    ledger = build_source_candidate_ledger(text)
    errors = verify_source_candidate_ledger(ledger, text)
    if errors:
        raise AssertionError(f"source ledger verification failed: {errors}")
    headers = [str(item) for item in expected["role_headers"]]
    layout = canonicalize_source_layout_v5(
        _empty_v5_layout(*(_alias_for(ledger, header) for header in headers)),
        ledger,
    )
    result = resolve_source_layout_v4(layout, ledger, classify_source_artifacts(ledger))
    if result.status != "verified" or result.layout is None:
        raise AssertionError(f"layout resolve failed: {result.status} {result.errors}")
    resolved = cast(SourceResumeLayoutReferencesV4, result.layout)
    document = materialize_source_document_v4(
        SourceResumeContentDocumentV4(candidates=ledger.candidates, layout=resolved)
    )
    if document is None:
        raise AssertionError("materialize_source_document_v4 returned None")
    return ledger, document


def source_backed_wrap_source_payload() -> tuple[dict[str, object], SourceResumeDocument]:
    """Return a verified v5 source payload with its matching materialized document."""
    text = load_wrap_resume_text()
    expected = load_expected_semantic()
    ledger = build_source_candidate_ledger(text)
    headers = [str(item) for item in expected["role_headers"]]
    layout = canonicalize_source_layout_v5(
        _empty_v5_layout(*(_alias_for(ledger, header) for header in headers)),
        ledger,
    )
    resolved = resolve_source_layout_v4(layout, ledger, classify_source_artifacts(ledger))
    if resolved.status != "verified" or resolved.layout is None:
        raise AssertionError(f"source layout resolve failed: {resolved.status} {resolved.errors}")
    document = SourceResumeContentDocumentV4(
        candidates=ledger.candidates,
        layout=cast(SourceResumeLayoutReferencesV4, resolved.layout),
    )
    payload: dict[str, object] = {
        "schema_version": "source-resume-v5",
        "document": document.model_dump(mode="json"),
        "layout_validation_status": "verified",
        "layout_validation_errors": [],
    }
    materialized = materialize_source_document_v4(document)
    if materialized is None:
        raise AssertionError("source fixture materialization failed")
    return payload, materialized


def source_backed_wrap_source_content() -> dict[str, object]:
    """Return a verified v5 source payload that the production service rehydrates."""
    payload, _ = source_backed_wrap_source_payload()
    return payload


def v5_source_row() -> SimpleNamespace:
    content = {
        "schema_version": "source-resume-v5",
        "document": {"candidates": [], "layout": {}, "artifacts": []},
        "layout_validation_status": "verified",
        "layout_validation_errors": [],
    }
    return SimpleNamespace(
        id=12,
        resume_id=13,
        validation_status="verified",
        content=content,
    )


def configure_v5_active_source(
    repository: Any,
    _source: SourceResumeDocument,
) -> SimpleNamespace:
    source_row = v5_source_row()
    repository.get_active_tailoring_source.return_value = (
        SimpleNamespace(
            id=8,
            parsed=ParsedProfile.model_validate({"skills": ["Python", "FastAPI"]}).model_dump(),
        ),
        source_row,
    )
    return source_row


def source_backed_wrap_content(
    source: SourceResumeDocument,
    source_row: Any,
    *,
    job_description: str = "Python backend platform engineer",
    suggested_first_bullet: str | None = None,
) -> TailoredResume:
    """Build a full-context snapshot from the sanitized wrap source."""
    # Keep suggested text source-identical so snapshot validation does not
    # depend on the LLM sanitizer. Distinct wording is covered in frontend
    # review-state tests.
    first_suggestion = suggested_first_bullet
    work_history: list[TailoredWorkItem] = []
    for role_index, source_role in enumerate(source.work_history):
        decisions = []
        for index, bullet in enumerate(source_role.highlights):
            suggested = (
                first_suggestion
                if first_suggestion is not None and role_index == 0 and index == 0
                else bullet.text
            )
            decisions.append(
                {
                    "source_id": bullet.source_id,
                    "source_index": index,
                    "canonical_order": bullet.order,
                    "original_text": bullet.text,
                    "suggested_text": suggested,
                    "final_text": suggested,
                    "recommended_included": True,
                    "final_included": True,
                    "recommendation_rank": index,
                    "output_rank": index,
                }
            )
        work_history.append(
            TailoredWorkItem(
                company=source_role.company.text if source_role.company is not None else None,
                title=source_role.title.text if source_role.title is not None else None,
                original_highlights=[bullet.text for bullet in source_role.highlights],
                tailored_highlights=[item["final_text"] for item in decisions],
                tailored_highlight_source_indices=list(range(len(decisions))),
                source_role_id=source_role.source_id,
                source_header=source_role.company.text if source_role.company is not None else None,
                role_context=[item.text for item in source_role.context_items],
                responsibilities_heading=(
                    source_role.responsibilities_heading.text
                    if source_role.responsibilities_heading is not None
                    else None
                ),
                selection_decisions=decisions,
            )
        )
    content = TailoredResume(
        job_id=11,
        professional_summary=None,
        full_name=None,
        email=None,
        phone=None,
        headline=None,
        skills=[],
        links={},
        projects=[],
        languages=[],
        licenses=[],
        education=[],
        certifications=[],
        work_history=work_history,
        canonical_sections=_canonical_sections(source),
        tailoring_snapshot_version="full-context-v1",
        review_evidence_version="review-evidence-v1",
        summary_evidence_ids=[],
        skill_order_source_ids=[],
        source_resume_id=source_row.resume_id,
        source_resume_content_id=source_row.id,
        source_schema_version="source-resume-v5",
        source_revision_sha256=source_revision_sha256(source_row.content),
    )
    for tailored_role, recommended in zip(
        content.work_history,
        select_source_bullets(source.work_history, job_description).recommended_ids,
        strict=True,
    ):
        for decision in tailored_role.selection_decisions:
            source_id = str(decision.source_id)
            decision.recommended_included = source_id in recommended
            decision.recommendation_rank = (
                recommended.index(source_id) if source_id in recommended else None
            )
        included = _required_included_indices(tailored_role.selection_decisions)
        tailored_role.tailored_highlight_source_indices = included
        tailored_role.tailored_highlights = [
            tailored_role.selection_decisions[index].final_text for index in included
        ]
        for index, decision in enumerate(tailored_role.selection_decisions):
            decision.output_rank = included.index(index) if index in included else None
    return content


def apply_review_choices(
    content: TailoredResume,
    *,
    restore_original: set[int] | None = None,
    omit: set[int] | None = None,
) -> TailoredResume:
    """Apply first-role review choices and rebuild recommended-first output order."""
    restore_original = restore_original or set()
    omit = omit or set()
    updated = content.model_copy(deep=True)
    role = updated.work_history[0]
    for index, decision in enumerate(role.selection_decisions):
        if index in restore_original:
            decision.final_text = decision.original_text
        if index in omit:
            decision.final_included = False
            decision.output_rank = None
    included = _required_included_indices(role.selection_decisions)
    role.tailored_highlight_source_indices = included
    role.tailored_highlights = [role.selection_decisions[index].final_text for index in included]
    for index, decision in enumerate(role.selection_decisions):
        decision.output_rank = included.index(index) if index in included else None
    return updated


def cli(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Extract normalized text from a tailored PDF or DOCX for semantic "
            "comparison. Ignores fonts, glyphs, and page breaks."
        )
    )
    parser.add_argument("paths", nargs="+", help="PDF, DOCX, or TXT files to extract")
    parser.add_argument(
        "--expected",
        type=Path,
        default=EXPECTED_SEMANTIC_PATH,
        help="expected_semantic.json used for token comparison",
    )
    args = parser.parse_args(argv)
    expected = json.loads(args.expected.read_text(encoding="utf-8"))
    failed = False
    for path_str in args.paths:
        path = Path(path_str)
        data = path.read_bytes()
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            surface = extract_pdf_text(data)
        elif suffix == ".docx":
            surface = extract_docx_text(data)
        else:
            surface = data.decode("utf-8")
        print(f"--- {path.name} ---")
        print(surface)
        try:
            assert_surface_matches_expected(surface, expected)
            print("SEMANTIC MATCH")
        except AssertionError as exc:
            failed = True
            print(f"SEMANTIC MISMATCH: {exc}")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(cli(sys.argv[1:]))
