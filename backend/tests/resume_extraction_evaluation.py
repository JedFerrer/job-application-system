"""Deterministic native PDF/DOCX extraction corpus support."""

from __future__ import annotations

import hashlib
import json
import re
import unicodedata
from collections.abc import Callable
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocxDocument

from app.infrastructure.resume_parser import ExtractedResumeText, extract_text_from_bytes

MANIFEST_PATH = Path(__file__).parent / "fixtures" / "resume_extraction" / "manifest.json"
_FIXTURE_ROOT = MANIFEST_PATH.parent
_SCHEMA_VERSION = "resume-extraction-eval-v1"
_ALLOWED_MANIFEST_FIELDS = {"schema_version", "provenance", "cases"}
_ALLOWED_CASE_FIELDS = {
    "id", "format", "fixture", "builder", "sha256", "byte_size",
    "expected_blocks", "expected_warning_codes", "expected_error", "interpretation",
}
_WARNING_CODES = {"short_extracted_text", "low_pdf_text_coverage", "likely_scanned_pdf"}
_PDF_FIXTURES = {"multipage.pdf", "two-column.pdf", "sparse-pages.pdf", "blank-pages.pdf"}
_TWO_COLUMN_INTERPRETATION = (
    "Pinned pdfplumber extractor order; not asserted as human reading order."
)


@dataclass(frozen=True)
class ExtractionCase:
    id: str
    format: str
    fixture: str | None
    builder: str | None
    sha256: str | None
    byte_size: int | None
    expected_blocks: tuple[str, ...] | None
    expected_warning_codes: tuple[str, ...] | None
    expected_error: str | None


def _docx_paragraphs_unicode(document: DocxDocument) -> None:
    document.add_paragraph("Nova Vale — fictional engineer")
    document.add_paragraph("Skills: Python; TypeScript; café résumé tooling")
    document.add_paragraph("- Built safe parsers (100% synthetic) — verify: § • №")


def _docx_paragraph_table_order(document: DocxDocument) -> None:
    document.add_paragraph("Before the fictional table")
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "Skill"
    table.cell(0, 1).text = "Python"
    table.cell(0, 2).text = ""
    table.cell(1, 0).text = ""
    table.cell(1, 1).text = "Atlas project"
    table.cell(1, 2).text = "example.invalid"
    document.add_paragraph("After the fictional table")


def _docx_multiple_tables(document: DocxDocument) -> None:
    first = document.add_table(rows=2, cols=2)
    first.cell(0, 0).text = "Contact"
    first.cell(0, 1).text = "fictional.invalid"
    second = document.add_table(rows=2, cols=2)
    second.cell(0, 0).text = "Project"
    cell = second.cell(0, 1)
    cell.text = "First paragraph"
    cell.add_paragraph("Second paragraph")
    cell.add_paragraph("Dense synthetic table evidence " + ("verified " * 80).strip())
    second.cell(1, 0).text = ""
    second.cell(1, 1).text = ""


_DOCX_BUILDERS: dict[str, Callable[[DocxDocument], None]] = {
    "paragraphs-unicode": _docx_paragraphs_unicode,
    "paragraph-table-order": _docx_paragraph_table_order,
    "multiple-tables": _docx_multiple_tables,
}


def build_docx(builder: str) -> bytes:
    try:
        build = _DOCX_BUILDERS[builder]
    except KeyError as exc:
        raise ValueError(f"unknown DOCX builder {builder!r}") from exc
    document = Document()
    build(document)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _required(mapping: dict[str, Any], field: str, case_label: str) -> Any:
    if field not in mapping:
        raise ValueError(f"{case_label} missing required field {field}")
    return mapping[field]


def _validate_expectations(
    item: dict[str, Any], case_id: str
) -> tuple[tuple[str, ...] | None, tuple[str, ...] | None, str | None]:
    has_success = "expected_blocks" in item or "expected_warning_codes" in item
    has_failure = "expected_error" in item
    if has_success == has_failure:
        raise ValueError(f"{case_id} must declare success or failure expectations")
    if has_failure:
        expected_error = item["expected_error"]
        if not isinstance(expected_error, str) or not expected_error:
            raise ValueError(f"{case_id}.expected_error must be a non-empty string")
        return None, None, expected_error
    blocks = item.get("expected_blocks")
    warnings = item.get("expected_warning_codes")
    if (
        not isinstance(blocks, list)
        or not blocks
        or not all(isinstance(x, str) and x for x in blocks)
    ):
        raise ValueError(f"{case_id}.expected_blocks must be non-empty strings")
    if not isinstance(warnings, list) or not all(isinstance(x, str) for x in warnings):
        raise ValueError(f"{case_id}.expected_warning_codes must be strings")
    if len(set(warnings)) != len(warnings):
        raise ValueError(f"{case_id}.expected_warning_codes contains duplicate warning codes")
    unknown = set(warnings) - _WARNING_CODES
    if unknown:
        raise ValueError(
            f"{case_id}.expected_warning_codes contains unknown code {sorted(unknown)}"
        )
    return tuple(blocks), tuple(warnings), None


def load_manifest(path: Path = MANIFEST_PATH) -> tuple[ExtractionCase, ...]:
    try:
        raw = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ValueError(f"malformed JSON manifest: {path}") from exc
    if not isinstance(raw, dict):
        raise ValueError("manifest must be an object")
    unknown_manifest = set(raw) - _ALLOWED_MANIFEST_FIELDS
    if unknown_manifest:
        raise ValueError(f"manifest unexpected fields: {sorted(unknown_manifest)}")
    if raw.get("schema_version") != _SCHEMA_VERSION:
        raise ValueError(f"unsupported schema version; expected {_SCHEMA_VERSION}")
    provenance = raw.get("provenance")
    if not isinstance(provenance, str) or not provenance.strip():
        raise ValueError("manifest missing provenance")
    cases_raw = raw.get("cases")
    if not isinstance(cases_raw, list) or len(cases_raw) != 7:
        raise ValueError("manifest must contain exactly 7 cases")

    cases: list[ExtractionCase] = []
    seen: set[str] = set()
    for index, item in enumerate(cases_raw):
        label = f"case[{index}]"
        if not isinstance(item, dict):
            raise ValueError(f"{label} must be an object")
        unknown = set(item) - _ALLOWED_CASE_FIELDS
        if unknown:
            raise ValueError(f"{label} unexpected fields: {sorted(unknown)}")
        case_id = _required(item, "id", label)
        if not isinstance(case_id, str) or not case_id:
            raise ValueError(f"{label}.id must be a non-empty string")
        if case_id in seen:
            raise ValueError(f"{case_id} duplicate id")
        seen.add(case_id)
        file_format = _required(item, "format", case_id)
        if file_format not in {"pdf", "docx"}:
            raise ValueError(f"{case_id} unsupported format {file_format!r}")
        if case_id == "pdf-two-column":
            interpretation = _required(item, "interpretation", case_id)
            if interpretation != _TWO_COLUMN_INTERPRETATION:
                raise ValueError(
                    f"{case_id}.interpretation must equal {_TWO_COLUMN_INTERPRETATION!r}"
                )
        elif "interpretation" in item:
            raise ValueError(f"{case_id} unexpected field interpretation")

        fixture = builder = sha256 = None
        byte_size = None
        if file_format == "pdf":
            if "builder" in item:
                raise ValueError(f"{case_id} unexpected fields for PDF: ['builder']")
            fixture = _required(item, "fixture", case_id)
            if not isinstance(fixture, str) or fixture not in _PDF_FIXTURES:
                raise ValueError(f"{case_id} unknown PDF fixture {fixture!r}")
            sha256 = _required(item, "sha256", case_id)
            byte_size = _required(item, "byte_size", case_id)
            if not isinstance(sha256, str) or not re.fullmatch(r"[0-9a-f]{64}", sha256):
                raise ValueError(f"{case_id}.sha256 must be a lowercase SHA-256 hex digest")
            if not isinstance(byte_size, int) or isinstance(byte_size, bool) or byte_size <= 0:
                raise ValueError(f"{case_id}.byte_size must be positive")
            fixture_path = _FIXTURE_ROOT / fixture
            if not fixture_path.is_file():
                raise ValueError(f"{case_id}.fixture file is missing")
            data = fixture_path.read_bytes()
            actual_size = len(data)
            if actual_size != byte_size:
                raise ValueError(
                    f"Case {case_id!r} fixture {fixture!r} byte size mismatch: "
                    f"expected {byte_size}, actual {actual_size}"
                )
            actual_sha256 = hashlib.sha256(data).hexdigest()
            if actual_sha256 != sha256:
                raise ValueError(
                    f"Case {case_id!r} fixture {fixture!r} SHA-256 mismatch: "
                    f"expected {sha256}, actual {actual_sha256}"
                )
        else:
            unexpected = {"fixture", "sha256", "byte_size"} & set(item)
            if unexpected:
                raise ValueError(f"{case_id} unexpected fields for DOCX: {sorted(unexpected)}")
            builder = _required(item, "builder", case_id)
            if not isinstance(builder, str) or builder not in _DOCX_BUILDERS:
                raise ValueError(f"{case_id} unknown DOCX builder {builder!r}")

        expected_blocks, expected_warning_codes, expected_error = _validate_expectations(
            item, case_id
        )
        cases.append(
            ExtractionCase(
                case_id,
                file_format,
                fixture,
                builder,
                sha256,
                byte_size,
                expected_blocks,
                expected_warning_codes,
                expected_error,
            )
        )

    if sum(case.format == "pdf" for case in cases) != 4:
        raise ValueError("manifest must contain exactly 4 PDF cases")
    if sum(case.format == "docx" for case in cases) != 3:
        raise ValueError("manifest must contain exactly 3 DOCX cases")
    return tuple(cases)


def normalize_blocks(text: str) -> tuple[str, ...]:
    normalized = unicodedata.normalize(
        "NFC", text.replace("\r\n", "\n").replace("\r", "\n")
    ).strip()
    return tuple(block for block in re.split(r"\n{2,}", normalized) if block) if normalized else ()


def evaluate_case(case: ExtractionCase) -> dict[str, Any]:
    if case.format == "pdf":
        assert case.fixture is not None
        content = (_FIXTURE_ROOT / case.fixture).read_bytes()
        filename = case.fixture
    else:
        assert case.builder is not None
        content = build_docx(case.builder)
        filename = f"{case.builder}.docx"
    result: dict[str, Any] = {
        "schema_version": _SCHEMA_VERSION,
        "case_id": case.id,
        "format": case.format,
        "expected_blocks": list(case.expected_blocks) if case.expected_blocks is not None else None,
        "actual_blocks": None,
        "expected_warning_codes": (
            list(case.expected_warning_codes)
            if case.expected_warning_codes is not None
            else None
        ),
        "actual_warning_codes": None,
        "expected_error": case.expected_error,
        "actual_error": None,
        "passed": False,
    }
    try:
        extracted = extract_text_from_bytes(content, filename, include_warnings=True)
        if not isinstance(extracted, ExtractedResumeText):
            raise TypeError("extraction seam returned an unexpected result")
        actual_blocks = normalize_blocks(extracted.text)
        actual_warnings = tuple(warning.code for warning in extracted.warnings)
        result["actual_blocks"] = list(actual_blocks)
        result["actual_warning_codes"] = list(actual_warnings)
        result["passed"] = (
            case.expected_error is None
            and actual_blocks == case.expected_blocks
            and actual_warnings == case.expected_warning_codes
        )
    except Exception as exc:
        result["actual_error"] = str(exc)
        result["passed"] = case.expected_error is not None and str(exc) == case.expected_error
    return result


def mismatch_diagnostic(result: dict[str, Any]) -> str:
    expected = result["expected_blocks"] or []
    actual = result["actual_blocks"] or []
    for index, (expected_block, actual_block) in enumerate(zip(expected, actual, strict=False)):
        if expected_block != actual_block:
            return (
                f"case {result['case_id']} block {index}: expected {expected_block!r}; "
                f"actual {actual_block!r}"
            )
    if len(expected) != len(actual):
        return (
            f"case {result['case_id']} block count: expected {len(expected)}; "
            f"actual {len(actual)}"
        )
    return f"case {result['case_id']} produced no block mismatch"


def evaluate_manifest() -> tuple[dict[str, Any], ...]:
    return tuple(evaluate_case(case) for case in load_manifest())
